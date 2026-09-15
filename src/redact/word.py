"""Redact a Word document in place, keeping its styles and automatic numbering.

Clause and schedule numbers aren't typed in the text: Word generates them from the
paragraph styles, so editing only the text of each run keeps them. A paragraph is kept
whole if it is a heading (all in capitals, such as "THE FIRST SCHEDULE") or the title
line that follows one of the headings named in the redaction settings. In a table, the
first column (the defined terms, the particulars' labels) is kept.
"""
from __future__ import annotations

import re
from pathlib import Path

from redact.text import TOKEN, Redactor


def _is_heading(text: str) -> bool:
    letters = [c for c in text if c.isalpha()]
    return bool(letters) and all(c.isupper() for c in letters)


def _redact_paragraph(paragraph, redactor: Redactor) -> None:
    runs = [r for r in paragraph.runs if r.text]
    text = "".join(r.text for r in runs)
    mask = redactor.kept(text)
    offset = 0
    for run in runs:
        start = offset
        offset += len(run.text)
        pieces, last = [], 0
        for m in TOKEN.finditer(run.text):
            s, e = start + m.start(), start + m.end()
            pieces += [run.text[last:m.start()], redactor.partial(m.group(), mask[s:e])]
            last = m.end()
        run.text = "".join(pieces) + run.text[last:]


def redact_docx(source: Path, target: Path, redactor: Redactor, titles_after: list[str] = ()) -> None:
    import docx

    document = docx.Document(source)
    after = [re.compile(p) for p in titles_after]
    keep_next = False
    for paragraph in document.paragraphs:
        text = " ".join(paragraph.text.split())
        if not text:
            continue
        if keep_next or _is_heading(text):
            keep_next = any(p.search(text) for p in after)
            continue
        _redact_paragraph(paragraph, redactor)
    for table in document.tables:
        for row in table.rows:
            for index, cell in enumerate(row.cells):
                if index == 0:
                    continue
                for paragraph in cell.paragraphs:
                    _redact_paragraph(paragraph, redactor)
    for section in document.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            for paragraph in part.paragraphs:
                _redact_paragraph(paragraph, redactor)
    core = document.core_properties
    core.author = core.last_modified_by = ""
    core.title = f"Redacted: {source.stem}"
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
