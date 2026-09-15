"""Word (.docx) documents -> line stream.

Word stores automatic numbering as list definitions, not as text: a clause that
prints as "7.12 The Lessee shall…" is stored as "The Lessee shall…"
in a Heading 2 paragraph. So the reader works out each paragraph's list label the way
Word does and puts it back in front of the text, and the parse rules then see the
document as it prints.

How the labels are computed:
- a paragraph's list comes from its own numbering properties, or failing that from its
  style (and the styles that style is based on);
- counters are kept per abstract list, so every w:num that shares an abstract list
  continues the same count;
- a w:num with a start override restarts its levels the first time it is used;
- counting a level resets every deeper level; empty paragraphs still count.

Tables are read separately (see docx_tables) and are not part of the line stream.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from .extract import Line

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _val(element, tag: str) -> str | None:
    found = element.find(W + tag) if element is not None else None
    return found.get(W + "val") if found is not None else None


@dataclass
class _Level:
    start: int = 1
    fmt: str = "decimal"
    text: str = "%1"


@dataclass
class Numbering:
    """The document's list definitions, and the running counters while reading."""

    abstract_of: dict[str, str] = field(default_factory=dict)  # numId -> abstractNumId
    levels: dict[str, dict[int, _Level]] = field(default_factory=dict)  # abstractNumId -> ilvl -> level
    overrides: dict[str, dict[int, int]] = field(default_factory=dict)  # numId -> ilvl -> start
    counters: dict[str, dict[int, int]] = field(default_factory=dict)  # abstractNumId -> ilvl -> value
    restarted: set[str] = field(default_factory=set)

    @classmethod
    def from_document(cls, document) -> "Numbering":
        numbering = cls()
        try:
            root = document.part.numbering_part.element
        except (KeyError, NotImplementedError):
            return numbering
        for abstract in root.findall(W + "abstractNum"):
            aid = abstract.get(W + "abstractNumId")
            numbering.levels[aid] = {
                int(level.get(W + "ilvl")): _Level(
                    int(_val(level, "start") or 1), _val(level, "numFmt") or "decimal", _val(level, "lvlText") or ""
                )
                for level in abstract.findall(W + "lvl")
            }
        for num in root.findall(W + "num"):
            nid = num.get(W + "numId")
            numbering.abstract_of[nid] = _val(num, "abstractNumId")
            numbering.overrides[nid] = {
                int(o.get(W + "ilvl")): int(_val(o, "startOverride"))
                for o in num.findall(W + "lvlOverride")
                if _val(o, "startOverride") is not None
            }
        return numbering

    def label(self, num_id: str, ilvl: int) -> str:
        aid = self.abstract_of.get(num_id)
        if aid is None or aid not in self.levels:
            return ""
        levels, counters = self.levels[aid], self.counters.setdefault(aid, {})
        if num_id not in self.restarted and self.overrides.get(num_id):
            for level, start in self.overrides[num_id].items():
                counters[level] = start - 1
                for deeper in [k for k in counters if k > level]:
                    del counters[deeper]
            self.restarted.add(num_id)
        level = levels.get(ilvl, _Level())
        counters[ilvl] = counters.get(ilvl, level.start - 1) + 1
        for deeper in [k for k in counters if k > ilvl]:
            del counters[deeper]
        if level.fmt == "none":
            return ""

        def part(m: re.Match) -> str:
            k = int(m.group(1)) - 1
            return str(counters.get(k, levels.get(k, _Level()).start))

        return re.sub(r"%(\d)", part, level.text).strip()


def _numbering_of(paragraph) -> tuple[str, int] | None:
    """(numId, ilvl) from the paragraph, or from its style chain."""
    ppr = paragraph._p.pPr
    num_id = ilvl = None
    if ppr is not None and ppr.numPr is not None:
        num_id = ppr.numPr.numId.val if ppr.numPr.numId is not None else None
        ilvl = ppr.numPr.ilvl.val if ppr.numPr.ilvl is not None else None
    style = paragraph.style
    while (num_id is None or ilvl is None) and style is not None:
        spr = style.element.pPr
        if spr is not None and spr.numPr is not None:
            if num_id is None and spr.numPr.numId is not None:
                num_id = spr.numPr.numId.val
            if ilvl is None and spr.numPr.ilvl is not None:
                ilvl = spr.numPr.ilvl.val
        style = style.base_style
    if num_id is None or str(num_id) == "0":
        return None
    return str(num_id), int(ilvl or 0)


def _is_bold(paragraph) -> bool:
    runs = [r for r in paragraph.runs if r.text.strip()]
    if not runs:
        return False
    style_bold = bool(paragraph.style.font.bold)
    return all((r.bold if r.bold is not None else style_bold) for r in runs)


def docx_lines(source: Path, cfg: dict) -> list[Line]:
    import docx

    document = docx.Document(source)
    numbering = Numbering.from_document(document)
    lines: list[Line] = []
    for paragraph in document.paragraphs:
        found = _numbering_of(paragraph)
        label = numbering.label(*found) if found else ""
        text = " ".join(paragraph.text.replace("\t", " ").split())
        if not text:
            continue  # counted above, but nothing to show
        text = f"{label} {text}" if label else text
        lines.append(Line(text, _is_bold(paragraph), None, None, None, paragraph.style.name))
    return lines


def docx_tables(source: Path) -> list[tuple[str, list[list[str]]]]:
    """Every table as (text of the paragraph before it, rows of cell text)."""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(source)
    tables, previous = [], ""
    for child in document.element.body.iterchildren():
        if child.tag == W + "p":
            text = Paragraph(child, document).text.strip()
            previous = text or previous
        elif child.tag == W + "tbl":
            rows = [[cell.text.strip() for cell in row.cells] for row in Table(child, document).rows]
            tables.append((previous, rows))
    return tables
